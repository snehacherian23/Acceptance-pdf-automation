"""
app/services/report_service.py
---------------------------------
Generates the per-job Excel processing report (processing_report.xlsx)
required by the project spec: "Export individual data in an Excel
format." One row per input item (Excel row, or signed document) --
success or failure -- so the report is a complete audit trail of the
whole batch, not just the successful ones.

Shared by both workflows: Workflow 1 (Excel rows -> documents) and
Workflow 2 (signed documents -> final PDFs) have different data
shapes, so `columns` is passed in per call rather than hardcoded --
avoids duplicating this whole file for a second report layout.

Uses openpyxl directly (already a project dependency) rather than
pandas, since this is a straightforward "list of dicts -> spreadsheet"
job with light formatting -- no need to pull pandas' heavier DataFrame
machinery in just for this.
"""

import os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

# Default column layout -- Workflow 1 (Excel row -> Acceptance Test document).
WORKFLOW1_COLUMNS = [
    ("source_file", "Source Excel File"),
    ("row_number", "Row Number"),
    ("customer_name", "Customer Name"),
    ("account_number", "Account Number"),
    ("property_id", "Property ID"),
    ("status", "Processing Status"),
    ("error_message", "Error Message"),
    ("output_document_name", "Output Document Name"),
    ("processed_at", "Processing Date & Time"),
]

# Workflow 2 (signed document + Consolidated Sheets -> final PDF).
WORKFLOW2_COLUMNS = [
    ("input_file", "Signed Document"),
    ("consolidated_count", "Consolidated Sheets Appended"),
    ("status", "Processing Status"),
    ("error_message", "Error Message"),
    ("output_file", "Output PDF"),
    ("processed_at", "Processing Date & Time"),
]

HEADER_FILL = PatternFill(start_color="4D363D", end_color="4D363D", fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True)
SUCCESS_FILL = PatternFill(start_color="E8F5EC", end_color="E8F5EC", fill_type="solid")
FAILED_FILL = PatternFill(start_color="FDEAEA", end_color="FDEAEA", fill_type="solid")


def generate_report(records, output_path, columns=None):
    """
    Build a processing report .xlsx from a list of per-item record
    dicts collected during processing (see job_tracker.add_record).

    Args:
        records: list of dicts with keys matching `columns`' first element.
        output_path: full path to write the .xlsx file to.
        columns: list of (record_key, header_label) tuples defining
                 the report's columns and their order. Defaults to
                 WORKFLOW1_COLUMNS for backward compatibility.

    Returns:
        output_path
    """
    columns = columns or WORKFLOW1_COLUMNS

    wb = Workbook()
    ws = wb.active
    ws.title = "Processing Report"

    # --- Header row ---
    for col_idx, (_, header) in enumerate(columns, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center")

    # --- Data rows ---
    for row_idx, record in enumerate(records, start=2):
        status = record.get("status", "")
        row_fill = SUCCESS_FILL if status == "Success" else FAILED_FILL

        for col_idx, (key, _) in enumerate(columns, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=record.get(key, ""))
            cell.fill = row_fill

    # --- Column widths (rough auto-fit based on header + a sane minimum) ---
    for col_idx, (key, header) in enumerate(columns, start=1):
        max_len = max(
            [len(str(header))] + [len(str(r.get(key, ""))) for r in records]
        ) if records else len(header)
        ws.column_dimensions[get_column_letter(col_idx)].width = min(max(max_len + 3, 14), 45)

    ws.freeze_panes = "A2"  # keep header visible while scrolling

    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    wb.save(output_path)
    return output_path
